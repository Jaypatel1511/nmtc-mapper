"""
Download and cache the CDFI Fund NMTC eligibility file.
Builds a lookup table of the FULL census-tract universe (85,395 tracts: 35,335
eligible + 50,060 ineligible), each carrying an explicit YES/NO LIC flag. A
tract ABSENT from this table is therefore unknown/indeterminate, not ineligible.
"""
import os
import re
import zipfile
import requests
import pandas as pd
from pathlib import Path

from nmtcmapper.exceptions import (
    EligibilityDataError, EligibilityDownloadError, EligibilityParseError,
    EligibilitySchemaError, EligibilityValueError,
    OZDataError, OZDownloadError, OZParseError,
    OZ2DataError, OZ2DownloadError, OZ2ParseError, OZ2SchemaError,
)
from nmtcmapper.data.schema import (
    CACHE_DIR, CDFI_FUND_LIC_URL_2020,
    ELIGIBILITY_XLSB_SHEET, ELIGIBILITY_XLSB_COLUMN_COUNT,
    ELIGIBILITY_XLSB_EXPECTED_HEADERS, ELIGIBILITY_MIN_ROWS,
    ELIGIBILITY_HEADER_SEARCH_ROWS,
    ELIGIBILITY_VALUE_BOUNDS, ELIGIBILITY_XLSB_VALUE_ALLOWLISTS,
    LIC_POVERTY_RATE_THRESHOLD,
    LIC_AMI_RATIO_METRO_THRESHOLD,
    LIC_AMI_RATIO_RURAL_THRESHOLD,
    SEVERE_POVERTY_THRESHOLD, SEVERE_AMI_THRESHOLD,
    SEVERE_UNEMPLOYMENT_MULTIPLIER, NATIONAL_UNEMPLOYMENT_RATE,
    DEEP_POVERTY_THRESHOLD, DEEP_AMI_THRESHOLD,
    DEEP_UNEMPLOYMENT_MULTIPLIER,
    OZ2_URL, OZ2_SHEET, OZ2_COLUMN_COUNT, OZ2_EXPECTED_HEADERS,
    OZ2_MIN_ROWS, OZ2_FLAG_ALLOWED, OZ2_TRACT_BINDING,
    OZ2_ELIGIBLE_LIC_COLUMN, OZ2_RURAL_STATUS_COLUMN,
)


def _normalize_header(value) -> str:
    """Collapse internal whitespace, strip, and casefold a header cell."""
    return re.sub(r"\s+", " ", str(value)).strip().casefold()


# Appended to every schema-drift message (0.4.2). This guard has now fired in
# the field once: the CDFI Fund re-published the eligibility file at the SAME URL
# in July 2026 with two renamed headers, and every live load failed until the
# pinned constants were updated. When that happens the exception is the only
# thing the user sees, so it has to carry the explanation and the remedy.
#
# Deliberately offers NO bypass. There is no safe one: the loader binds columns
# positionally, so continuing past a header it does not recognize means reading
# eligibility out of an unverified slot — a wrong NMTC answer that looks right.
_DRIFT_REMEDY = (
    "\n\nWhy this happens: the CDFI Fund re-publishes the eligibility file IN "
    "PLACE, at the same URL, without changing the filename — so a new layout can "
    "arrive under a file this package already believed it understood. nmtc-mapper "
    "pins the exact header strings on purpose, and fails here rather than parse a "
    "changed file against stale positions."
    "\n\nWhat to do: upgrade nmtc-mapper — a release that recognizes the new "
    "layout may already exist (pip install --upgrade nmtc-mapper). If you are "
    "already on the latest version, please report this, quoting the mismatch "
    "above: https://github.com/Jaypatel1511/nmtc-mapper/issues"
    "\n\nIf you reached this from a cached copy, a fresh download has already "
    "been attempted automatically and also failed to match, so the live file "
    "genuinely differs from what this release pins."
    "\n\nThere is no supported way to bypass this check. The verdicts it would "
    "let through are not trustworthy."
)


def _validate_xlsb_header(header_vals: list) -> None:
    """Fail loud (EligibilitySchemaError) BEFORE any row is parsed if the live
    data-sheet structure does not match the expected CDFI Fund layout.

    The loader binds columns positionally, so a wrong column count or a
    renamed/re-ordered column at a bound index must be caught here — otherwise a
    poverty rate could be read out of the MFI slot and every verdict would be
    silently wrong."""
    n = len(header_vals)
    if n != ELIGIBILITY_XLSB_COLUMN_COUNT:
        raise EligibilitySchemaError(
            f"eligibility workbook header has {n} columns, expected "
            f"{ELIGIBILITY_XLSB_COLUMN_COUNT}. The column layout has changed — "
            f"the positional bind can no longer be trusted."
            + _DRIFT_REMEDY
        )
    for idx, expected in ELIGIBILITY_XLSB_EXPECTED_HEADERS.items():
        actual = header_vals[idx] if idx < n else None
        if _normalize_header(actual) != _normalize_header(expected):
            raise EligibilitySchemaError(
                f"eligibility workbook header mismatch at column index {idx}: "
                f"expected {expected!r}, got {actual!r}. The loader binds columns "
                f"positionally, so a renamed/re-ordered column would be read "
                f"against the wrong field."
                + _DRIFT_REMEDY
            )


def _check_value_bounds(field: str, value, row_index: int) -> None:
    """Raise EligibilityValueError if a numeric value is implausible (Fix 6).

    None (an 'NA' cell) is a legitimate null — it is NEVER bounds-checked."""
    if value is None:
        return
    lo, hi = ELIGIBILITY_VALUE_BOUNDS[field]
    if not (lo <= value <= hi):
        raise EligibilityValueError(
            f"{field} out of plausible bounds at data row {row_index}: value "
            f"{value!r} not in [{lo}, {hi}]. (ami_ratio is stored as a FRACTION "
            f"~0.9; a value near percent scale ~91 signals an upstream 100x "
            f"scale flip that would break every AMI comparison.)"
        )


def _checked_cell(col_index: int, raw, row_index: int) -> str:
    """Normalize a categorical cell and reject any value outside its allowlist (0.5.0).

    Returns the `.strip().upper()` form so the caller compares against a value
    this function has already vouched for.

    The header guard pins header STRINGS; it cannot see a change to cell
    VOCABULARY (schema.ELIGIBILITY_XLSB_VALUE_ALLOWLISTS). Without this check the
    `== "YES"` tests on columns 2/13/14/15 map every unrecognized value to
    `False` — `'Y'` parses to `False` — which is a fabricated negative on the LIC
    verdict and on both distress flags. The `!= "METRO"` test on column 1 fails
    the other way, mapping anything unrecognized to `True`. Both are silent.
    Raise instead: this loader binds columns positionally and there is no safe
    way to continue past a value it does not recognize.
    """
    label, allowed = ELIGIBILITY_XLSB_VALUE_ALLOWLISTS[col_index]
    normalized = str(raw).strip().upper()
    if normalized not in allowed:
        raise EligibilitySchemaError(
            f"eligibility workbook column {label} carries an unrecognized value at "
            f"data row {row_index}: {raw!r} (normalizes to {normalized!r}). "
            f"Expected one of {sorted(allowed)}."
            f"\n\nThe header guard cannot catch this: it pins header strings, not "
            f"cell vocabularies, so a re-publish that leaves every header "
            f"byte-identical and rewrites a cell value passes it completely. "
            f"Parsing on would silently map this value to a verdict this package "
            f"cannot support."
            + _DRIFT_REMEDY
        )
    return normalized


def get_cache_dir() -> Path:
    path = Path(CACHE_DIR)
    path.mkdir(parents=True, exist_ok=True)
    return path


def _cache_path(filename: str) -> Path:
    return get_cache_dir() / filename


# The cache filename never changes — and neither does the CDFI Fund's URL when
# they re-publish, which is why a warm cache can hold a superseded layout.
#
# 0.6.1: renamed from `..._2016_2020.xlsb`. The Fund now publishes an .xlsx and
# the extension here records that — but the loader NEVER trusts it: the parser
# is chosen by sniffing the ZIP (`_sniff_workbook_format`), so xlsb bytes under
# this name would still be read by pyxlsb. The rename also means an upgrade
# from <= 0.6.0 does a plain cold download rather than a self-heal; the old
# `.xlsb` file, if present, is simply unused and may be deleted.
ELIGIBILITY_CACHE_FILENAME = "NMTC_LIC_Eligibility_2016_2020.xlsx"


def _eligibility_cache_path() -> Path:
    return _cache_path(ELIGIBILITY_CACHE_FILENAME)


# Both workbooks this package downloads (.xlsb and .xlsx) are OOXML, i.e. ZIP
# containers, so every legitimate body starts with the ZIP local-file-header
# magic. An HTML page does not.
_ZIP_MAGIC = b"PK\x03\x04"

# Small floor purely against a zero-length or stub body that happened to start
# with the magic. The real content check is the header/row-count validation
# downstream; this is only about never installing obvious garbage as the cache.
_MIN_WORKBOOK_BYTES = 4096


def _validated_workbook_bytes(tmp: Path, url: str) -> None:
    """Raise if `tmp` is not plausibly a workbook. Call BEFORE tmp.replace(path).

    An HTTP 200 carrying an HTML body is the shape a CDN/origin stack serves
    during maintenance, and ``raise_for_status()`` is blind to it: the status is
    200, the stream writes fine, and before 0.4.2 the HTML went straight to
    ``tmp.replace(path)`` and destroyed a good 4.8 MB cache. Every later run then
    served those bytes from cache, failed to parse them, and never re-downloaded
    — the package stayed broken until the user found and deleted the cache
    directory by hand. The download must therefore prove the body is a workbook
    before it is allowed anywhere near the cache path.
    """
    size = tmp.stat().st_size if tmp.exists() else 0
    head = tmp.read_bytes()[:len(_ZIP_MAGIC)] if size else b""
    if head != _ZIP_MAGIC or size < _MIN_WORKBOOK_BYTES:
        preview = tmp.read_bytes()[:80] if size else b""
        raise EligibilityDownloadError(
            f"Download from {url} returned {size:,} bytes that are not an "
            f"Excel workbook (expected a ZIP/OOXML container starting "
            f"{_ZIP_MAGIC!r}, got {head!r}). This is the signature of an HTTP 200 "
            f"carrying an error or maintenance page rather than the file. "
            f"The cached copy has NOT been replaced. First bytes: {preview!r}"
        )


def download_eligibility_file(force: bool = False) -> Path:
    path = _eligibility_cache_path()
    if path.exists() and not force:
        print(f"Using cached eligibility file: {path}")
        return path
    print("Downloading NMTC eligibility file from CDFI Fund...")
    # Stream to a .part temp and atomically rename on success, so a mid-stream
    # failure can never leave a truncated file at the final cache path (which
    # would make every later run parse-fail on poisoned cache).
    tmp = path.with_suffix(path.suffix + ".part")
    try:
        response = requests.get(CDFI_FUND_LIC_URL_2020, stream=True, timeout=120)
        response.raise_for_status()
        with open(tmp, "wb") as f:
            for chunk in response.iter_content(chunk_size=8192):
                f.write(chunk)
        # Nothing unvalidated may reach the cache path: replacing a good 4.8 MB
        # workbook with an HTML maintenance page is not recoverable from here.
        try:
            _validated_workbook_bytes(tmp, CDFI_FUND_LIC_URL_2020)
        except EligibilityDownloadError:
            tmp.unlink(missing_ok=True)
            raise
        tmp.replace(path)
        print(f"Saved to {path}")
        return path
    except requests.exceptions.HTTPError as e:
        tmp.unlink(missing_ok=True)
        status = getattr(e.response, "status_code", None)
        if status == 404:
            reason = "not found (404) — the CDFI Fund file path may have moved"
        elif status == 403:
            reason = "access blocked (403 Forbidden)"
        else:
            reason = f"HTTP {status}"
        raise EligibilityDownloadError(
            f"Failed to download NMTC eligibility file from "
            f"{CDFI_FUND_LIC_URL_2020}: {reason}"
        ) from e
    except requests.exceptions.RequestException as e:
        tmp.unlink(missing_ok=True)
        raise EligibilityDownloadError(
            f"Failed to download NMTC eligibility file from "
            f"{CDFI_FUND_LIC_URL_2020}: connection/DNS/timeout error "
            f"({type(e).__name__}: {e})"
        ) from e


def load_eligibility_table(force: bool = False) -> pd.DataFrame:
    """Load the CDFI Fund eligibility table, preferring the cached copy.

    Stale-cache self-heal (0.4.2): the CDFI Fund re-publishes this file IN PLACE
    under an unchanged filename, so a cache warmed by an earlier release can hold
    a superseded layout. Validating the new pins against those old bytes would
    raise EligibilitySchemaError and tell the user to upgrade — advice they had
    just taken. So a schema mismatch on a CACHED file re-downloads once and
    re-validates. This does NOT weaken the guard: the fresh file is checked just
    as strictly, and a genuine upstream divergence still raises (once, not in a
    loop). An explicit force=True already fetched fresh bytes, so it never
    retries.

    The self-heal covers an UNREADABLE cached file as well as a stale-layout one.
    EligibilityParseError is a sibling of EligibilitySchemaError, not a subclass,
    so catching only the latter left a cached file that cannot be parsed at all
    outside the heal: it raised on every run and never triggered a re-download,
    and the user stayed broken until they deleted ~/.nmtcmapper/cache by hand.
    That is exactly the state a truncated write or (before this release) an
    HTTP-200 HTML body left behind, and the cache is the one thing the package
    can fix for itself. Both are caught on a CACHED file only, and both still
    retry at most once.

    The at-most-once property is structural, not a counter: the handler calls the
    private ``_load_eligibility_table(force=True)``, not this wrapper, so the
    retry cannot itself re-enter the heal.
    """
    served_from_cache = (not force) and _eligibility_cache_path().exists()
    try:
        return _load_eligibility_table(force=force)
    except (EligibilitySchemaError, EligibilityParseError) as e:
        if not served_from_cache:
            raise
        if isinstance(e, EligibilitySchemaError):
            print(
                "Cached eligibility file does not match the expected CDFI Fund "
                "layout. The Fund re-publishes in place, so the cached copy may "
                "be superseded — re-downloading once and re-validating..."
            )
        else:
            print(
                "Cached eligibility file could not be read at all — the cached "
                "bytes are corrupt or are not a workbook. Discarding them and "
                "re-downloading once..."
            )
        return _load_eligibility_table(force=True)


def _load_eligibility_table(force: bool = False) -> pd.DataFrame:
    path = download_eligibility_file(force=force)
    if path is None or not path.exists():
        # download_eligibility_file now raises on any download failure, so a
        # missing file here means the download step was skipped and no local
        # copy exists. Fail loud — never substitute demo data (F2).
        raise EligibilityDownloadError(
            f"No eligibility file available at {_eligibility_cache_path()} and no "
            f"download was performed."
        )
    print(f"Loading eligibility table from {path}...")
    try:
        # 0.6.1: ONE positional parser, TWO container readers, chosen by the
        # bytes. (0.5.0 deleted a generic-workbook `else` branch here as
        # structurally unreachable dead code keyed on `path.suffix`; this is
        # not that branch coming back — nothing here consults the suffix.)
        return _load_eligibility_workbook(path)
    except EligibilityDataError:
        raise
    except ImportError:
        # Missing optional engine (e.g. pyxlsb) — its own message is actionable
        # and is a dependency problem, not a corrupt-file problem. Don't mask it.
        raise
    except Exception as e:
        raise EligibilityParseError(
            f"Failed to parse eligibility file {path}: {type(e).__name__}: {e}"
            f"\n\nThose bytes are not a readable workbook. If they came from the "
            f"cache, a fresh download has already been attempted automatically "
            f"and the file still would not parse, so the problem is upstream or "
            f"on the network path — not a stale cache. You can delete "
            f"{path.parent} to start clean; nothing in it is user data."
        ) from e


# Both containers this loader accepts are OOXML ZIPs; the ONE member that
# differs is the workbook part. This is what the dispatch reads — never the
# URL's extension, never the cache filename.
_XLSB_WORKBOOK_PART = "xl/workbook.bin"
_XLSX_WORKBOOK_PART = "xl/workbook.xml"


def _sniff_workbook_format(path: Path) -> str:
    """Return "xlsb" or "xlsx" from the bytes at `path`; raise a named error otherwise.

    0.6.1. Before this, the parser was chosen by fiat (pyxlsb, always) and the
    September-2026 .xlsx died inside pyxlsb with a bare
    ``KeyError: "There is no item named 'xl/_rels/workbook.bin.rels' in the
    archive"`` — not even a package exception. An unreadable file must raise a
    named error that quotes what was actually found, so the person reading it
    can tell an HTML error page from a new container from a truncated write.

    Raises EligibilityParseError, which the cached-file self-heal in
    ``load_eligibility_table`` treats as "discard and re-download once".
    """
    size = path.stat().st_size if path.exists() else 0
    head = path.read_bytes()[:16] if size else b""
    if not head.startswith(_ZIP_MAGIC):
        raise EligibilityParseError(
            f"eligibility file {path} is not an Excel workbook: expected a "
            f"ZIP/OOXML container starting {_ZIP_MAGIC!r}, got {size:,} bytes "
            f"starting {head!r}. An HTML error page, an empty file and a "
            f"truncated write all look like this."
        )
    try:
        with zipfile.ZipFile(path) as zf:
            names = zf.namelist()
    except zipfile.BadZipFile as e:
        raise EligibilityParseError(
            f"eligibility file {path} starts with the ZIP magic {_ZIP_MAGIC!r} "
            f"but is not a readable ZIP ({e}); first bytes {head!r}. The bytes "
            f"are corrupt or truncated."
        ) from e
    if _XLSB_WORKBOOK_PART in names:
        return "xlsb"
    if _XLSX_WORKBOOK_PART in names:
        return "xlsx"
    raise EligibilityParseError(
        f"eligibility file {path} is a ZIP but not a spreadsheet this package "
        f"can read: it has neither {_XLSB_WORKBOOK_PART!r} (.xlsb) nor "
        f"{_XLSX_WORKBOOK_PART!r} (.xlsx). First bytes {head!r}; members found: "
        f"{names[:8]}{' ...' if len(names) > 8 else ''}. If the CDFI Fund has "
        f"moved to a third container format, this package needs a release."
    )


def _iter_xlsb_rows(path: Path):
    """Yield each row of the data sheet as a list of cell values (pyxlsb)."""
    try:
        import pyxlsb
    except ImportError:
        raise ImportError(
            "pyxlsb is required to read the CDFI Fund .xlsb file. "
            "Install it with: pip install pyxlsb"
        )
    with pyxlsb.open_workbook(str(path)) as wb:
        with wb.get_sheet(ELIGIBILITY_XLSB_SHEET) as sheet:
            for row in sheet.rows():
                yield [c.v for c in row]


def _iter_xlsx_rows(path: Path):
    """Yield each row of the data sheet as a list of cell values (openpyxl).

    read_only + data_only: the September-2026 file is a 60 MB sheet XML with a
    3.5 MB calcChain; the streaming reader is what keeps this within the memory
    the .xlsb path used, and data_only means a formula cell yields its cached
    value rather than the formula text."""
    import openpyxl
    # Hand openpyxl an open FILE, not the path: given a path string it refuses
    # any extension outside its own list before looking at a byte, and the
    # cache filename is a module constant that says nothing about the bytes
    # in it. The dispatch is on the container, all the way down.
    with open(path, "rb") as fh:
        wb = openpyxl.load_workbook(fh, read_only=True, data_only=True)
        try:
            if ELIGIBILITY_XLSB_SHEET not in wb.sheetnames:
                raise EligibilityParseError(
                    f"eligibility file {path} has no sheet named "
                    f"{ELIGIBILITY_XLSB_SHEET!r} (sheets present: {wb.sheetnames})."
                )
            for row in wb[ELIGIBILITY_XLSB_SHEET].iter_rows(values_only=True):
                yield list(row)
        finally:
            wb.close()


def _load_xlsb_table(path: Path) -> pd.DataFrame:
    """The CDFI Fund eligibility workbook as an .xlsb (Aug 2025 - Sep 2026)."""
    return _parse_eligibility_rows(_iter_xlsb_rows(path))


def _load_xlsx_table(path: Path) -> pd.DataFrame:
    """The CDFI Fund eligibility workbook as an .xlsx (Sep 2026 -)."""
    return _parse_eligibility_rows(_iter_xlsx_rows(path))


def _load_eligibility_workbook(path: Path) -> pd.DataFrame:
    """Dispatch on the container the bytes actually are. See _sniff_workbook_format."""
    fmt = _sniff_workbook_format(path)
    if fmt == "xlsb":
        return _load_xlsb_table(path)
    return _load_xlsx_table(path)


def _parse_eligibility_rows(rows) -> pd.DataFrame:
    """Parse the CDFI Fund data sheet from a row iterator (either container).

    Column layout (0-indexed), confirmed against the live file (September-2026
    .xlsx; identical positions in the Aug-2025b/July-2026 .xlsb):
      0  GEOID, 1 Metro/Non-metro, 2 LIC eligible (YES/NO),
      3  Poverty rate %, 5 MFI ratio (decimal), 7 Unemployment rate %,
     13  High Migration Rural County tract (YES/NO), 14 Severe distress (YES/NO),
     15  Deep distress (YES/NO)

    `rows` yields one list of cell VALUES per sheet row, in sheet order, starting
    at the sheet's first row. The header row is located, not assumed: it is the
    first row with a non-blank column 0 within ELIGIBILITY_HEADER_SEARCH_ROWS
    (the September-2026 .xlsx carries a banner row above it; the .xlsb did not).

    0.6.1 NOTE ON COLUMN N. The paragraphs below were written against the
    July-2026 file, where column N held 1,422 YES and was titled "... Low-Income
    Community Census Tract". The September-2026 file retitles it "... Census
    Tract for Deep Distress" and holds 1,318 YES — the 104 dropped are poverty-
    route LICs with MFI above 85%, so the column now carries the INCOME-route
    determination only. Every check below still holds on the new file (the
    1,318 are all column-C YES, all non-metro, all MFI <= 85% or NA), the OR is
    still a no-op, and 0 of 85,395 verdicts moved. The figures in the
    paragraphs are left as measured on the file they name. See schema.py.

    THE LIC VERDICT IS COLUMN 2 **OR** COLUMN 13 (0.4.2).
    ------------------------------------------------------------------
    Through 0.4.1 the verdict was column 2 alone. That under-reported 168 tracts
    for the whole life of the .xlsb path (v0.3.1 .. v0.4.1): under the Aug-2025b
    layout, column 2 carried only the poverty and <=80%-MFI criteria, and the
    third statutory LIC route — a tract in a high migration rural county with MFI
    at or below 85% of the applicable area MFI, 26 U.S.C. 45D(e)(5), added by
    section 223 of the American Jobs Creation Act of 2004 (P.L. 108-357) — was
    published only in column 13. The July-2026 re-publish widened column 2 to
    absorb column 13, which is why 0.4.2 reads correctly from column 2 alone.
    That is upstream's choice, not a property of this package, and nothing in the
    loader would detect the Fund separating the two columns again: the header
    guard, the row-count floor and the value bounds all pass on such a file and
    the 168 silently go back to "ineligible". OR-ing column 13 in makes the
    verdict follow the statute rather than upstream's current formatting.

    Column 13 is an LIC DETERMINATION, not bare county membership. Establishing
    that is what makes the OR safe: if column 13 meant only "sits in a high
    migration rural county", OR-ing it would grant LIC to HMR tracts whose MFI
    exceeds 85% — the mirror of the defect being fixed. Three independent checks
    on the live file say it is a determination:
      * Its 1,422 YES tracts sit in 437 counties that together hold 1,883 tracts.
        The other 461 are column-13 NO, so the flag is a strict subset of county
        membership. None of those 461 reaches 20% poverty and only 2 fall in the
        <=85% band (see below).
      * Every one of the 1,422 satisfies a statutory LIC prong: 757 on poverty
        >=20%, 1,084 on MFI <=80%, and the remaining 168 in the (80%, 85%] band
        that 45D(e)(5) opens. Zero fail all prongs.
      * The Fund's own widened column 2 is exactly `col4 OR col6 OR col13`
        (poverty flag, MFI<=80% flag, this column) with zero mismatches across
        all 85,395 rows. OR-ing column 13 reproduces the Fund's own rule.
    The Aug-2025b workbook, whose extra "High migration tracts" sheet the
    re-publish dropped, states the same in prose: its county sub-table is headed
    "High Migration Counties (only Low-Income Community census tracts within
    counties are eligible)", and its tract sub-table lists exactly the 168.

    On the live July-2026 file this OR is a NO-OP: all 1,422 column-13 YES rows
    are already column-2 YES, so the eligible count is 35,335 with and without
    it. It is a floor under the verdict, not a change to it.
    """

    def _num(v):
        # bool is an int subclass — exclude it so a stray YES/NO never divides.
        return v if isinstance(v, (int, float)) and not isinstance(v, bool) else None

    def _blank_cell(v):
        return v is None or (isinstance(v, str) and not v.strip())

    records = []
    header_seen = False
    for i, vals in enumerate(rows):
        if not header_seen:
            if not vals or _blank_cell(vals[0]):
                # A pre-header row (the September-2026 banner). Bounded.
                if i + 1 >= ELIGIBILITY_HEADER_SEARCH_ROWS:
                    raise EligibilitySchemaError(
                        f"no header row found in the first "
                        f"{ELIGIBILITY_HEADER_SEARCH_ROWS} rows of sheet "
                        f"{ELIGIBILITY_XLSB_SHEET!r}: every one has a blank "
                        f"column 0. The layout has changed."
                        + _DRIFT_REMEDY
                    )
                continue
            # Validate structure BEFORE trusting any positional bind.
            _validate_xlsb_header(vals)
            header_seen = True
            continue
        if not vals or not vals[0]:
            continue

        geoid        = str(vals[0]).strip().zfill(11)
        # Every categorical cell goes through the value allowlist first
        # (0.5.0). `_checked_cell` raises on anything unrecognized rather
        # than letting `!= "METRO"` drift True or `== "YES"` drift False.
        non_metro    = _checked_cell(1, vals[1], i) == "NON-METRO"
        poverty_rate = (_num(vals[3]) / 100) if _num(vals[3]) is not None else None
        ami_ratio    = _num(vals[5])
        unemp_rate   = (_num(vals[7]) / 100) if _num(vals[7]) is not None else None
        high_migr    = _checked_cell(13, vals[13], i) == "YES"
        severe       = _checked_cell(14, vals[14], i) == "YES"
        deep         = _checked_cell(15, vals[15], i) == "YES"

        # The LIC verdict is column C **OR** column N, never column C
        # alone. See the docstring above for why column N is an LIC
        # determination and not bare "in a high migration rural county",
        # and why this is a no-op today.
        lic_elig     = (_checked_cell(2, vals[2], i) == "YES") or high_migr

        # Value plausibility (Fix 6) — on the STORED value; None passes.
        _check_value_bounds("poverty_rate", poverty_rate, i)
        _check_value_bounds("ami_ratio", ami_ratio, i)
        _check_value_bounds("unemployment_rate", unemp_rate, i)

        if deep:
            dlevel = "deep"
        elif severe:
            dlevel = "severe"
        elif lic_elig:
            dlevel = "lic"
        else:
            dlevel = "ineligible"

        records.append({
            "tract_id":              geoid,
            "nmtc_eligible":         lic_elig,
            "distress_level":        dlevel,
            "poverty_rate":          poverty_rate,
            "ami_ratio":             ami_ratio,
            "unemployment_rate":     unemp_rate,
            "is_non_metro":          non_metro,
            "is_high_migration_rural": high_migr,
            "severe_distress":       severe,
            "deep_distress":         deep,
        })

    if not header_seen:
        raise EligibilitySchemaError(
            f"sheet {ELIGIBILITY_XLSB_SHEET!r} has no rows at all — no header "
            f"could be validated." + _DRIFT_REMEDY
        )

    # Row-count floor: a degenerate/near-empty parse must raise, not yield an
    # (almost) empty table that would silently mark real tracts "not found".
    if len(records) < ELIGIBILITY_MIN_ROWS:
        raise EligibilitySchemaError(
            f"eligibility workbook parsed only {len(records)} data rows, below the "
            f"floor of {ELIGIBILITY_MIN_ROWS} (live file has 85,395). This is a "
            f"degenerate parse, not a usable table."
        )

    df = pd.DataFrame(records).set_index("tract_id")
    print(f"Eligibility table loaded: {len(df):,} census tracts")
    return df


def _compute_eligibility(df: pd.DataFrame) -> pd.DataFrame:
    """Apply the NMTC LIC / distress rules to a metrics frame.

    BACKS `load_sample_table()` ONLY. The official `.xlsb` path never calls this:
    it reads the Fund's own published columns C/N/O/P. Through 0.4.3 a second
    call site existed on a generic-workbook branch; that branch was structurally
    unreachable and 0.5.0 deleted it.

    Three structural defects were corrected in 0.5.0, all over-inclusive:

    1. NO `AND LIC` CONJUNCTION on severe/deep. The Fund's own column headers read
       `Severe distress=LIC AND (...)` / `Deep distress=LIC AND (...)`: distress
       is a tier WITHIN eligibility, never a route into it. Poverty >= 30% implies
       poverty >= 20% and MFI <= 60% implies MFI <= 80%, so those two prongs
       cannot fire outside LIC — but unemployment carries no LIC implication.

       MEASURED ON THE LIVE 85,395 ROWS, AGAINST TWO NAMED BASELINES. Both are
       correct measurements and they are NOT the same quantity, so neither number
       may be quoted without saying which "NOT LIC" it is scored against:

         * NOT LIC per the Fund's PUBLISHED column C — the criterion baseline:
           5,197 severe / 751 deep.
         * NOT LIC per THIS RULE'S OWN `nmtc_eligible` output — the experienced
           baseline, what a 0.4.3 caller actually saw, because they read both
           fields off the same frame: 5,063 severe / 733 deep. Read off
           `distress_level` instead of the boolean columns it is the same 5,063:
           deep is a subset of severe here, and `distress_label` tests deep, then
           severe, and only then `nmtc_eligible`, so the label short-circuits
           before the LIC test and partitions that one population into 4,330
           "severe" + 733 "deep".

       The two baselines reconcile exactly through defect (2): 134 of the 5,197
       (18 of the 751) are rows this rule itself called LIC while the Fund did
       not, all 134 inside the 932 it granted LIC on non-metro status alone.
       5,197 - 134 = 5,063 and 751 - 18 = 733.

       Under either baseline every one of them is carried by the unemployment
       prong alone — poverty 0, MFI 0. deep is a subset of severe under this rule
       (every deep prong is strictly tighter than its severe counterpart, and NaN
       fails both), so the union IS the severe count in both columns.

       After the fix: 0 under BOTH baselines.

    2. `is_non_metro` STOOD IN FOR THE HIGH-MIGRATION-RURAL 85% BAND.
       §45D(e)(1)(B) sets the income test at 80% for EVERY tract; the 85% figure
       comes from §45D(e)(5)(A) alone, which attaches the substitution to
       paragraph (1)(B)(i) — the NON-METROPOLITAN branch — and nothing else. And
       §45D(e)(5)(B) defines "high migration rural county" by out-migration
       alone, with no rurality test and no metropolitan test, so a METROPOLITAN
       county can satisfy it; when it does, its tracts are governed by (1)(B)(ii),
       which the substitution never touches. The band therefore requires
       `hmr & ~metro`. On the current file all 1,422 HMR tracts are non-metro, so
       the conjunct is redundant as an empirical property of ONE PUBLISHED FILE,
       not as a logical one — which is exactly why it is written out. The shipped
       rule granted LIC to 932 tracts on non-metro status alone.

    3. `>=` vs `>` ON THE DISTRESS POVERTY PRONGS. The Fund publishes strictly
       greater (`Poverty>30%`, `Poverty>40%`; April 2025 Compliance FAQ Q32), and
       the boundary population decides it: of the LIC tracts at exactly 30.0%
       qualifying on poverty alone the Fund published severe=NO for all 21, and at
       exactly 40.0% deep=NO for all 13. THE LIC PRONG STAYS `>=` — §45D(e)(1)(A)
       says a poverty rate "of at least 20 percent".
    """
    pr = df.get("poverty_rate", pd.Series(dtype=float))
    ami = df.get("ami_ratio", pd.Series(dtype=float))
    unemp = df.get("unemployment_rate", pd.Series(dtype=float))
    non_metro = df.get("is_non_metro", pd.Series(False, index=df.index))
    high_migr = df.get("is_high_migration_rural", pd.Series(False, index=df.index))

    # LIC poverty prong is `>=` and MUST STAY `>=` — 45D(e)(1)(A), "at least".
    poverty_lic = pr >= LIC_POVERTY_RATE_THRESHOLD
    # 80% for every tract, PLUS the 85% band for high-migration-rural tracts that
    # are also non-metropolitan. See defect (2) above.
    ami_lic = (
        (ami <= LIC_AMI_RATIO_METRO_THRESHOLD) |
        (high_migr & non_metro & (ami <= LIC_AMI_RATIO_RURAL_THRESHOLD))
    )
    df["nmtc_eligible"] = poverty_lic | ami_lic
    lic = df["nmtc_eligible"]

    # Distress poverty prongs are STRICTLY greater, and both tiers are AND-ed
    # with LIC. AND-ing here — rather than at the columns' consumers — is what
    # also repairs `distress_label` below, which short-circuits on the distress
    # columns before it ever consults nmtc_eligible.
    sev_poverty = pr > SEVERE_POVERTY_THRESHOLD
    sev_ami = ami <= SEVERE_AMI_THRESHOLD
    sev_unemp = unemp >= (NATIONAL_UNEMPLOYMENT_RATE * SEVERE_UNEMPLOYMENT_MULTIPLIER)
    df["severe_distress"] = lic & (sev_poverty | sev_ami | sev_unemp)

    deep_poverty = pr > DEEP_POVERTY_THRESHOLD
    deep_ami = ami <= DEEP_AMI_THRESHOLD
    deep_unemp = unemp >= (NATIONAL_UNEMPLOYMENT_RATE * DEEP_UNEMPLOYMENT_MULTIPLIER)
    df["deep_distress"] = lic & (deep_poverty | deep_ami | deep_unemp)

    def distress_label(row):
        if row.get("deep_distress"):
            return "deep"
        elif row.get("severe_distress"):
            return "severe"
        elif row.get("nmtc_eligible"):
            return "lic"
        return "ineligible"

    df["distress_level"] = df.apply(distress_label, axis=1)
    return df


def load_sample_table() -> pd.DataFrame:
    """Build the 12-tract synthetic demo table — an EXPLICIT opt-in only.

    WARNING: this is 12 synthetic-vintage tracts for demos, examples, and tests.
    It is NEVER valid for a real NMTC eligibility answer. Before 0.3.4 this table
    was substituted silently whenever a download or parse failed, fabricating
    'ineligible' results for real tracts; that path is gone. Use it only through
    ``load_sample_table()`` or ``NMTCMapper.from_sample()`` when you knowingly
    want demo data with ``data_source == "sample"``.
    """
    # 0.5.0 dropped the trailing `native_area` element with the
    # `is_nmtc_native_area` field: it was `False` on all twelve rows and no source
    # this package loads could ever have made it True.
    sample_tracts = [
        ("17031840100", 0.38, 0.55, 0.12, False, False),
        ("17031839100", 0.42, 0.48, 0.15, False, False),
        ("17031010100", 0.18, 0.92, 0.04, False, False),
        ("36061015900", 0.35, 0.60, 0.11, False, False),
        ("36061019100", 0.28, 0.72, 0.09, False, False),
        ("36047052200", 0.14, 0.88, 0.05, False, False),
        ("26163518300", 0.45, 0.45, 0.18, False, False),
        ("26163520100", 0.32, 0.62, 0.13, False, False),
        ("13121010400", 0.29, 0.68, 0.10, False, False),
        ("48113010900", 0.22, 0.78, 0.07, False, False),
        ("17019000100", 0.15, 0.95, 0.03, True,  True),
        ("26001010100", 0.18, 0.88, 0.06, True,  False),
    ]
    rows = []
    for tid, pr, ami, unemp, non_metro, high_migration in sample_tracts:
        rows.append({
            "tract_id": tid,
            "state": tid[:2],
            "poverty_rate": pr,
            "ami_ratio": ami,
            "unemployment_rate": unemp,
            "is_non_metro": non_metro,
            "is_high_migration_rural": high_migration,
        })
    df = pd.DataFrame(rows)
    df = _compute_eligibility(df)
    df = df.set_index("tract_id")
    return df


# Backwards-compatible alias: the examples/ notebook imports the private name.
# Kept intentionally so that existing imports keep working after the F4 rename.
_build_sample_table = load_sample_table


# ── Opportunity Zone lookup ───────────────────────────────────────────────────

def load_opportunity_zones(force: bool = False) -> set:
    """
    Load the set of Opportunity Zone census tract IDs.
    Returns a set of 11-digit FIPS codes designated as QOZs.

    Source: CDFI Fund designated-qozs.12.14.18.xlsx (8,764 tracts).
    Sheet "QOZs 14Jun", header on row 5 (index 4),
    tract FIPS in column "Census Tract Number".
    Raises OZDownloadError / OZParseError on any failure — never falls back.

    VINTAGE CAVEAT: these are 2010-based census tracts (OZs were designated in
    2018 on 2010 geography). The geocoder returns 2020 tracts, so ~16% of these
    GEOIDs (1,408 / 8,764) have no row in the 2020-basis eligibility table.

    0.5.0: THAT IS WHY ``is_opportunity_zone`` IS NEVER ``False``. A vintage miss
    and a genuine non-designation are the same observation without a crosswalk,
    so a non-match returns ``None`` — see ``EligibilityResult`` and
    ``opportunity_zone_status``. Membership is tested against this set directly
    and NOT against ``tract_found``, so a retired 2010 GEOID that is designated
    still returns a correct ``True``.
    """
    from nmtcmapper.data.schema import OZ_URL_2018

    filename = "QOZ_Designated_2018.xlsx"
    path = _cache_path(filename)

    if not path.exists() or force:
        # Same atomic .part-then-rename pattern as download_eligibility_file —
        # a mid-stream failure must never poison the final cache path.
        tmp = path.with_suffix(path.suffix + ".part")
        try:
            print("Downloading Opportunity Zone tract list...")
            response = requests.get(OZ_URL_2018, stream=True, timeout=60)
            response.raise_for_status()
            with open(tmp, "wb") as f:
                for chunk in response.iter_content(chunk_size=8192):
                    f.write(chunk)
            # Same pre-replace guard as the eligibility download: an HTTP-200
            # HTML body must never overwrite a good cached workbook. (The OZ
            # path has no self-heal, so poisoning here would be permanent —
            # which is why nothing unvalidated may reach tmp.replace().)
            try:
                _validated_workbook_bytes(tmp, OZ_URL_2018)
            except EligibilityDownloadError as e:
                tmp.unlink(missing_ok=True)
                raise OZDownloadError(str(e)) from e
            tmp.replace(path)
            print(f"Saved to {path}")
        except requests.exceptions.HTTPError as e:
            tmp.unlink(missing_ok=True)
            status = getattr(e.response, "status_code", None)
            if status == 404:
                reason = "not found (404) — the OZ file path may have moved"
            elif status == 403:
                reason = "access blocked (403 Forbidden)"
            else:
                reason = f"HTTP {status}"
            raise OZDownloadError(
                f"Failed to download Opportunity Zone file from {OZ_URL_2018}: {reason}"
            ) from e
        except requests.exceptions.RequestException as e:
            tmp.unlink(missing_ok=True)
            raise OZDownloadError(
                f"Failed to download Opportunity Zone file from {OZ_URL_2018}: "
                f"connection/DNS/timeout error ({type(e).__name__}: {e})"
            ) from e

    try:
        df = pd.read_excel(
            path,
            sheet_name="QOZs 14Jun",
            header=4,       # row 5 is the column header row
            dtype=str,
        )
        # Normalize column names: collapse whitespace/newlines, uppercase
        df.columns = df.columns.str.replace(r"\s+", " ", regex=True).str.strip().str.upper()
        for col in ["CENSUS TRACT NUMBER", "GEOID", "CENSUS_TRACT", "TRACT_ID", "TRACT"]:
            if col in df.columns:
                tracts = set(df[col].dropna().str.strip().str.zfill(11).tolist())
                print(f"Loaded {len(tracts):,} Opportunity Zone tracts")
                return tracts
    except OZDataError:
        raise
    except Exception as e:
        raise OZParseError(
            f"Failed to parse Opportunity Zone file {path}: {type(e).__name__}: {e}"
        ) from e

    # File parsed, but none of the known tract columns were present — fail loud
    # rather than degrade to the 6-tract sample (F3).
    raise OZParseError(
        f"Opportunity Zone file {path} parsed but no tract column was found "
        f"(looked for CENSUS TRACT NUMBER / GEOID / CENSUS_TRACT / TRACT_ID / TRACT)."
    )


def _sample_oz_tracts() -> set:
    """6 known OZ tracts — EXPLICIT opt-in only (used by NMTCMapper.from_sample()).

    As of 0.3.4 this is never reached from a download/parse failure path; OZ
    failures raise OZDownloadError / OZParseError instead of degrading to this set.
    """
    return {
        "17031840100",  # Chicago South Side
        "17031839100",  # Chicago West Side
        "26163518300",  # Detroit
        "36061015900",  # NYC Bronx
        "13121010400",  # Atlanta
        "48113010900",  # Dallas
    }


# ── OZ 2.0 nomination eligibility (0.6.0) ─────────────────────────────────────
#
# A SECOND SOURCE ON A SECOND SCHEME. Everything below is deliberately parallel
# to — and never shared with — the NMTC path above. See docs/oz2-methodology.md.
#
# WHY THERE IS NO SHARED, PARAMETERISED "LIC" HELPER between the two paths, and
# why there must never be one: OZ 2.0 and NMTC are structurally different tests.
# § 1400Z-1(c)(1) no longer cross-references § 45D(e) at all —
#
#     income threshold      NMTC 80%            OZ 2.0 70%
#     metro benchmark       GREATER of state    metro MFI only
#                           or metro MFI
#     non-metro benchmark   statewide MFI       statewide MFI
#     poverty path          >=20%, standalone   >=20% AND MFI <=125%
#
# — two independent tightenings, not one. Adding a `threshold=0.70` argument to
# anything that also serves NMTC would be the defect this note exists to prevent.
#
# IN PRACTICE THE CONSTRAINT DOES NOT BIND HERE, and saying so is more useful than
# an abstraction built to satisfy it: this loader READS Treasury's published
# `eligible_lic` column. It does not compute the OZ 2.0 LIC test at all, so there
# is no threshold to share and no helper to parameterise. The constraint is
# recorded because it binds hard on anyone who later tries to *compute* this test.

OZ2_CACHE_FILENAME = "OZ2_Eligible_LIC_Tracts_Treasury.xlsx"


def _oz2_cache_path() -> Path:
    return _cache_path(OZ2_CACHE_FILENAME)


def download_oz2_file(force: bool = False) -> Path:
    """Download Treasury's OZ 2.0 data-transparency file, atomically."""
    path = _oz2_cache_path()
    if path.exists() and not force:
        print(f"Using cached OZ 2.0 file: {path}")
        return path
    print("Downloading OZ 2.0 eligibility file from Treasury...")
    tmp = path.with_suffix(path.suffix + ".part")
    try:
        response = requests.get(OZ2_URL, stream=True, timeout=120)
        response.raise_for_status()
        with open(tmp, "wb") as f:
            for chunk in response.iter_content(chunk_size=8192):
                f.write(chunk)
        # Same pre-replace guard as both existing download paths: an HTTP 200
        # carrying an HTML maintenance page must never reach the cache.
        try:
            _validated_workbook_bytes(tmp, OZ2_URL)
        except EligibilityDownloadError as e:
            tmp.unlink(missing_ok=True)
            raise OZ2DownloadError(str(e)) from e
        tmp.replace(path)
        print(f"Saved to {path}")
        return path
    except requests.exceptions.HTTPError as e:
        tmp.unlink(missing_ok=True)
        status = getattr(e.response, "status_code", None)
        if status == 404:
            reason = ("not found (404) — Treasury moved this file once already; "
                      "oz-tracker's URL for it is a 404 for exactly this reason")
        elif status == 403:
            reason = "access blocked (403 Forbidden)"
        else:
            reason = f"HTTP {status}"
        raise OZ2DownloadError(
            f"Failed to download the OZ 2.0 eligibility file from {OZ2_URL}: {reason}"
        ) from e
    except requests.exceptions.RequestException as e:
        tmp.unlink(missing_ok=True)
        raise OZ2DownloadError(
            f"Failed to download the OZ 2.0 eligibility file from {OZ2_URL}: "
            f"connection/DNS/timeout error ({type(e).__name__}: {e})"
        ) from e


def _validate_oz2_header(header_vals: list) -> None:
    """Exact-match all twelve headers, in order. Raises OZ2SchemaError."""
    if len(header_vals) != OZ2_COLUMN_COUNT:
        raise OZ2SchemaError(
            f"OZ 2.0 sheet {OZ2_SHEET!r} has {len(header_vals)} columns, expected "
            f"{OZ2_COLUMN_COUNT}. Treasury has changed the file's shape; every "
            f"column binding below is now unverified."
        )
    got = tuple(_normalize_header(v) for v in header_vals)
    want = tuple(_normalize_header(v) for v in OZ2_EXPECTED_HEADERS)
    if got != want:
        diffs = [
            f"    {i}: expected {w!r}\n       live     {g!r}"
            for i, (w, g) in enumerate(zip(want, got)) if w != g
        ]
        raise OZ2SchemaError(
            "OZ 2.0 file header mismatch — Treasury has re-published this file "
            "with a different layout.\n" + "\n".join(diffs)
        )


def _blank(v) -> bool:
    """True for a Treasury blank: an empty cell or a whitespace-only string."""
    return v is None or (isinstance(v, str) and not v.strip())


def load_oz2_table(force: bool = False) -> pd.DataFrame:
    """Load Treasury's OZ 2.0 table, indexed by 11-digit GEOID.

    Columns:
      ``oz2_nomination_eligible``  bool  — Treasury's ``eligible_lic``, 1/0
      ``oz2_rural_area_qoz``       Optional[bool] — ``rural_status``, RESTRICTED
                                   to ``eligible_lic == 1``; None otherwise
      ``oz2_inputs_missing``       bool  — both poverty_rate AND mfi_ratio blank

    THE RURAL RESTRICTION IS DELIBERATE AND IS NOT A CONVENIENCE. Treasury
    populates ``rural_status`` for all 85,529 rows, but its rural methodology
    documents the determination over ELIGIBLE tracts only ("...which census
    tracts eligible to be nominated as 2027 QOZs ... are comprised entirely of a
    rural area"). A populated column is not a published determination. Exposing
    the other 20,377 rural-flagged-but-ineligible rows would be the same
    inference the methodology forbids for the Appendix — "it has a lot of rows,
    therefore it is authoritative".

    ``oz2_inputs_missing`` HAS NO ANALOGUE UPSTREAM AND IS NOT IN THE METHODOLOGY.
    It was found by execution while building this release, and it matters because
    Treasury's 0 is not one fact. See ``EligibilityResult.oz2_nomination_status``.
    """
    path = download_oz2_file(force=force)
    try:
        import openpyxl
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        try:
            if OZ2_SHEET not in wb.sheetnames:
                raise OZ2ParseError(
                    f"OZ 2.0 file {path} has no sheet named {OZ2_SHEET!r} "
                    f"(sheets present: {wb.sheetnames})."
                )
            rows = wb[OZ2_SHEET].iter_rows(values_only=True)
            header = list(next(rows))
            _validate_oz2_header(header)
            idx = {name: i for i, name in enumerate(OZ2_EXPECTED_HEADERS)}
            g_i = idx[OZ2_TRACT_BINDING.table_geoid_header]
            lic_i, rur_i = idx[OZ2_ELIGIBLE_LIC_COLUMN], idx[OZ2_RURAL_STATUS_COLUMN]
            pov_i, mfr_i = idx["poverty_rate"], idx["mfi_ratio"]

            geoids, elig, rural, missing = [], [], [], []
            for n, row in enumerate(rows, start=1):
                if row[g_i] is None:
                    continue
                # Treasury stores states 01-09 as zero-padded STRINGS and 10+ as
                # INTs (15,662 / 69,867 live). Both normalise to the same 11-digit
                # key under the loader's existing rule; without the zfill every
                # tract in states 01-09 would silently miss.
                raw = str(row[g_i]).strip()
                if raw.endswith(".0"):
                    raw = raw[:-2]
                geoid = raw.zfill(11)
                for col_i, label in ((lic_i, OZ2_ELIGIBLE_LIC_COLUMN),
                                     (rur_i, OZ2_RURAL_STATUS_COLUMN)):
                    val = row[col_i]
                    if val not in OZ2_FLAG_ALLOWED:
                        raise OZ2SchemaError(
                            f"OZ 2.0 column {label!r} carries {val!r} at data row "
                            f"{n} (GEOID {geoid}), which is outside the allowed "
                            f"{sorted(OZ2_FLAG_ALLOWED)}. A value this guard does "
                            f"not recognise would parse falsy and become a "
                            f"FABRICATED NEGATIVE."
                        )
                is_elig = bool(row[lic_i])
                geoids.append(geoid)
                elig.append(is_elig)
                # Restricted here, at the single point of construction, so no
                # consumer can reach an unrestricted rural flag at all.
                rural.append(bool(row[rur_i]) if is_elig else None)
                missing.append(_blank(row[pov_i]) and _blank(row[mfr_i]))
        finally:
            wb.close()
    except OZ2DataError:
        raise
    except ImportError:
        raise
    except Exception as e:
        raise OZ2ParseError(
            f"Failed to parse the OZ 2.0 file {path}: {type(e).__name__}: {e}"
        ) from e

    if len(geoids) < OZ2_MIN_ROWS:
        raise OZ2SchemaError(
            f"OZ 2.0 file {path} yielded only {len(geoids):,} rows (floor is "
            f"{OZ2_MIN_ROWS:,}). A degenerate parse must raise, never produce a "
            f"table in which every real tract reads as absent."
        )
    # THE CHECK TractVintage CANNOT PERFORM. Derived from the table's own keys and
    # asserted against the declaration — not read back off the declaration.
    OZ2_TRACT_BINDING.validate_table_scheme(geoids)

    # NOTE: plain lists, never pd.Series here. A Series carries its own
    # RangeIndex, and DataFrame(index=...) REINDEXES it against the GEOIDs
    # instead of labelling it — which silently turns every rural value into NaN.
    # Caught by execution while building this release, not by review.
    df = pd.DataFrame(
        {
            "oz2_nomination_eligible": elig,
            "oz2_rural_area_qoz": rural,
            "oz2_inputs_missing": missing,
        },
        index=pd.Index(geoids, name="tract_id"),
    )
    df["oz2_rural_area_qoz"] = df["oz2_rural_area_qoz"].astype(object)
    if df.index.has_duplicates:
        dupes = df.index[df.index.duplicated()].unique().tolist()[:5]
        raise OZ2SchemaError(
            f"OZ 2.0 file {path} carries duplicate GEOIDs (e.g. {dupes}). The "
            f"key is not unique, so no lookup against it is well-defined."
        )
    print(f"OZ 2.0 tracts loaded: {len(df):,}")
    return df


def _sample_oz2_table() -> pd.DataFrame:
    """Synthetic OZ 2.0 rows for ``NMTCMapper.from_sample()`` — NEVER a real answer.

    Deliberately keyed on the COG scheme so the sample cannot be joined against
    the NMTC sample's legacy CT keys by accident, and carries at least one row of
    every answer the real table can produce, including the no-inputs row.
    """
    rows = [
        # geoid,          eligible, rural_status, inputs_missing
        ("09110990100",   True,  True,  False),   # eligible + rural (CT, COG scheme)
        ("17031840100",   True,  False, False),   # eligible, not rural
        ("36061015900",   False, False, False),   # measured ineligible
        ("48113010900",   False, True,  False),   # ineligible but rural-flagged -> None
        ("06037990100",   False, False, True),    # published 0 with NO inputs at all
    ]
    df = pd.DataFrame(
        {
            "oz2_nomination_eligible": [r[1] for r in rows],
            "oz2_rural_area_qoz": [r[2] if r[1] else None for r in rows],
            "oz2_inputs_missing": [r[3] for r in rows],
        },
        index=pd.Index([r[0] for r in rows], name="tract_id"),
    )
    df["oz2_rural_area_qoz"] = df["oz2_rural_area_qoz"].astype(object)
    return df
